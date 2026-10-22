"""Build backlog.json, the Markdown user-stories doc and the XLSX backlog from the story data."""
import json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

import meta, stories_a, stories_b, stories_c

import stories_d
if not any(e == "LOC" for e, _ in meta.EPICS):
    meta.EPICS += stories_d.NEW_EPICS
meta.SCREENS.update(stories_d.NEW_SCREENS)
STORIES = stories_a.STORIES + stories_b.STORIES + stories_c.STORIES
for _s in STORIES:
    for _sc in stories_d.SCREEN_ADDITIONS.get(_s["id"], []):
        if _sc not in _s["screens"]:
            _s["screens"].append(_sc)
STORIES = STORIES + stories_d.STORIES
EPIC_NAME = dict(meta.EPICS)
EPIC_ORDER = [e for e, _ in meta.EPICS]
MS = {m["id"]: m for m in meta.MILESTONES}
MS_ORDER = [m["id"] for m in meta.MILESTONES]

def sprint_key(s):
    sp = s["sprint"]
    return int(sp[1:]) if sp else 99

json.dump({"personas": meta.PERSONAS, "milestones": meta.MILESTONES, "epics": meta.EPICS,
           "screens": {k: list(v) for k, v in meta.SCREENS.items()}, "stories": STORIES},
          open("backlog.json", "w"), indent=1, ensure_ascii=False)

# ---------------- Markdown ----------------
import os
ISSUES_MD = json.load(open("issue_map.json")) if os.path.exists("issue_map.json") else {}
L = []
w = L.append
w("# Trailbuddies: Personas and User Stories")
w("### Whole product, from MVP to post-launch\n")
w("*Prepared 21 October 2026 for Steve Gaita. Companion to `Trailbuddies-concept-report.md` and `Trailbuddies-development-plan.md`. "
  "The same backlog is in `Trailbuddies-backlog.xlsx` and as GitHub issues in the `trailbuddies` repo; story IDs match across all three.*\n")
w("---\n")
w("## How to read this\n")
w("- **Story format:** *As [persona], I want [goal], so that [benefit].* Acceptance criteria are written as **Given / When / Then** scenarios that can become automated tests.")
w("- **Priority (MoSCoW):** Must = the milestone fails without it · Should = important, can slip one sprint · Could = nice to have.")
w("- **Size:** XS (1 pt) · S (2) · M (3) · L (5) · XL (8). One point is roughly half a day of Claude-assisted work for a solo developer.")
w("- **Milestone and sprint** follow the development plan: " + " · ".join(f"**{m['id']}** {m['title'].split(' — ')[1]} ({m['sprints']})" for m in meta.MILESTONES) + ".")
w("- **Screens** refer to the wireframe IDs in the clickable prototype (P = parent phone, C = child tablet, T = teacher desktop, A = older learner).\n")

# summary
w("## Backlog at a glance\n")
cnt = collections.Counter(s["milestone"] for s in STORIES)
pts = collections.Counter()
for s in STORIES: pts[s["milestone"]] += meta.SIZE_POINTS[s["size"]]
w("| Milestone | Stories | Must | Should | Could | Points |")
w("|---|---|---|---|---|---|")
for m in MS_ORDER:
    pr = collections.Counter(s["priority"] for s in STORIES if s["milestone"] == m)
    w(f"| {MS[m]['title']} | {cnt[m]} | {pr['Must']} | {pr['Should']} | {pr['Could']} | {pts[m]} |")
w(f"| **Total** | **{len(STORIES)}** | **{sum(1 for s in STORIES if s['priority']=='Must')}** | **{sum(1 for s in STORIES if s['priority']=='Should')}** | **{sum(1 for s in STORIES if s['priority']=='Could')}** | **{sum(pts.values())}** |\n")
w("> **Capacity check:** the MVP holds " + str(pts['MVP']) + " points across 8 sprints (~" + str(round(pts['MVP']/8)) +
  " points a sprint). That is ambitious for one developer, even with Claude. If velocity after Sprint 2 is lower, move the MVP's *Should* and *Could* stories into Phase 2 first; the *Must* stories are the beta.\n")

w("| Epic | Stories | Points | Milestones |")
w("|---|---|---|---|")
for e in EPIC_ORDER:
    ss = [s for s in STORIES if s["epic"] == e]
    ms = sorted({s["milestone"] for s in ss}, key=MS_ORDER.index)
    w(f"| [{e}: {EPIC_NAME[e]}](#{e.lower()}-{EPIC_NAME[e].lower().replace(' & ', '--').replace(' ', '-').replace(':', '').replace('(', '').replace(')', '')}) | {len(ss)} | {sum(meta.SIZE_POINTS[s['size']] for s in ss)} | {', '.join(ms)} |")
w("")
w("---\n")

w("## Personas\n")
for p in meta.PERSONAS:
    w(f"### {p['name']} ({p['role']}{', age ' + p['age'] if p['age'] not in ('—',) else ''})")
    w(p["summary"] + "\n")
    w(f"- **Needs:** {p['needs']}")
    w(f"- **Frustrations:** {p['frustrations']}\n")
w("---\n")

w("## Stories by epic\n")
for e in EPIC_ORDER:
    ss = sorted([s for s in STORIES if s["epic"] == e], key=lambda s: (MS_ORDER.index(s["milestone"]), sprint_key(s), s["id"]))
    w(f"### {e}: {EPIC_NAME[e]}\n")
    for s in ss:
        w(f"#### {s['id']} · {s['title']}")
        meta_line = f"`{s['priority']}` · `{s['size']} ({meta.SIZE_POINTS[s['size']]} pt)` · `{s['milestone']}{' / ' + s['sprint'] if s['sprint'] else ''}`"
        if s["screens"]:
            meta_line += " · Screens: " + ", ".join(f"{sc} {meta.SCREENS[sc][0]}" for sc in s["screens"])
        if s["id"] in ISSUES_MD: meta_line += f" · [GitHub #{ISSUES_MD[s['id']]}](https://github.com/SGGaita/trailbuddies/issues/{ISSUES_MD[s['id']]})"
        w(meta_line + "\n")
        w(f"**{s['story']}**\n")
        for i, a in enumerate(s["ac"], 1):
            w(f"- *Scenario {i}:* **Given** {a['given']}, **when** {a['when']}, **then** {a['then']}.")
        w("")
w("---\n")

w("## Screen → story traceability\n")
w("| Screen | Device | Stories |")
w("|---|---|---|")
for sc, (title, dev, _) in meta.SCREENS.items():
    ids = [s["id"] for s in STORIES if sc in s["screens"]]
    w(f"| {sc} {title} | {dev} | {', '.join(ids)} |")
w("")
w("## Sprint plan (MVP)\n")
w("| Sprint | Stories | Points |")
w("|---|---|---|")
for n in range(0, 8):
    sp = f"S{n}"
    ss = [s for s in STORIES if s["sprint"] == sp]
    w(f"| {sp} | {', '.join(s['id'] for s in ss)} | {sum(meta.SIZE_POINTS[s['size']] for s in ss)} |")
w("")
open("Trailbuddies-user-stories.md", "w").write("\n".join(L))

# ---------------- XLSX ----------------
F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F3A5F")
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D0D0D0")

wb = Workbook()

# README
ws = wb.active; ws.title = "README"
rows = [
    ("Trailbuddies backlog", None),
    ("Prepared 21 Oct 2026 for Steve Gaita. Same story IDs as the Markdown doc and the GitHub issues.", None),
    ("", None),
    ("Sheet", "What it holds"),
    ("Backlog", "One row per user story. Edit Status, Sprint, Owner and GitHub # here; Points is calculated from Size."),
    ("Summary", "Counts and points by milestone, priority, status and epic; all formulas, so it updates as you edit Backlog."),
    ("Personas", "Who the stories are written for."),
    ("Screens", "Wireframe screen IDs used in the Screens column."),
    ("Lists", "Allowed values for Priority, Size (with points), Milestone and Status."),
    ("", None),
    ("Cells you edit", "Backlog: Status, Sprint, Owner, GitHub #, Notes (yellow headers). Everything else comes from the story source."),
    ("Size → points", "XS=1, S=2, M=3, L=5, XL=8 (one point ≈ half a day of Claude-assisted work). Set in Lists."),
    ("Example row", "MVP-ready stories look like ONB-02: Must · M · MVP · S2 · Status 'To do'."),
]
for r in rows: ws.append(list(r))
ws["A1"].font = Font(name=F, size=16, bold=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = body; c.alignment = wrap
for c in ws[4]: c.font = bold
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 100

# Lists
wl = wb.create_sheet("Lists")
wl.append(["Size", "Points", "", "Priority", "", "Milestone", "Title", "Due", "", "Status"])
sizes = list(meta.SIZE_POINTS.items())
statuses = ["To do", "In progress", "In review", "Done", "Blocked", "Won't do"]
for i in range(max(len(sizes), len(meta.MILESTONES), len(statuses), 3)):
    row = [""] * 10
    if i < len(sizes): row[0], row[1] = sizes[i]
    if i < 3: row[3] = ["Must", "Should", "Could"][i]
    if i < len(meta.MILESTONES): row[5], row[6], row[7] = meta.MILESTONES[i]["id"], meta.MILESTONES[i]["title"], meta.MILESTONES[i]["due"]
    if i < len(statuses): row[9] = statuses[i]
    wl.append(row)
for c in wl[1]: c.font = hdr_font; c.fill = hdr_fill
for row in wl.iter_rows(min_row=2):
    for c in row: c.font = body
for col, wdt in zip("ABCDEFGHIJ", [8, 8, 3, 10, 3, 10, 44, 12, 3, 14]): wl.column_dimensions[col].width = wdt
n_sizes = len(sizes)

# Backlog
wbk = wb.create_sheet("Backlog", 1)
cols = ["ID", "Epic", "Epic name", "Title", "Persona / role", "User story", "Acceptance criteria (Given / When / Then)",
        "Priority", "Size", "Points", "Milestone", "Sprint", "Screens", "Status", "Owner", "GitHub #", "Notes"]
wbk.append(cols)
ordered = sorted(STORIES, key=lambda s: (MS_ORDER.index(s["milestone"]), sprint_key(s), EPIC_ORDER.index(s["epic"]), s["id"]))
import os
ISSUES = json.load(open("issue_map.json")) if os.path.exists("issue_map.json") else {}
for i, s in enumerate(ordered, start=2):
    ac = "\n".join(f"{n}. GIVEN {a['given']}\n   WHEN {a['when']}\n   THEN {a['then']}" for n, a in enumerate(s["ac"], 1))
    wbk.append([s["id"], s["epic"], EPIC_NAME[s["epic"]], s["title"], s["role"], s["story"], ac,
                s["priority"], s["size"], f"=IFERROR(INDEX(Lists!$B$2:$B${n_sizes+1},MATCH(I{i},Lists!$A$2:$A${n_sizes+1},0)),0)",
                s["milestone"], s["sprint"], ", ".join(s["screens"]), "To do", "Steve", ISSUES.get(s["id"], ""), s["notes"]])
last = len(ordered) + 1
for c in wbk[1]:
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
for col in ("N", "L", "O", "P", "Q"):
    wbk[f"{col}1"].fill = PatternFill("solid", fgColor="B8860B")
widths = [9, 7, 22, 30, 20, 50, 70, 9, 6, 7, 10, 7, 14, 12, 10, 9, 20]
for idx, wdt in enumerate(widths, 1): wbk.column_dimensions[get_column_letter(idx)].width = wdt
for row in wbk.iter_rows(min_row=2, max_row=last):
    for c in row: c.font = body; c.alignment = wrap; c.border = Border(bottom=thin)
wbk.freeze_panes = "E2"
tab = Table(displayName="Backlog", ref=f"A1:Q{last}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
wbk.add_table(tab)
for col, lst in (("H", "Lists!$D$2:$D$4"), ("I", f"Lists!$A$2:$A${n_sizes+1}"), ("K", f"Lists!$F$2:$F${len(meta.MILESTONES)+1}"), ("N", f"Lists!$J$2:$J${len(statuses)+1}")):
    dv = DataValidation(type="list", formula1=f"={lst}", allow_blank=False)
    wbk.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")

# Summary
wsu = wb.create_sheet("Summary", 2)
wsu.append(["Backlog summary (all formulas: updates as the Backlog sheet changes)"])
wsu["A1"].font = Font(name=F, size=14, bold=True)
wsu.append([])
wsu.append(["Milestone", "Title", "Stories", "Must", "Should", "Could", "Points", "Done (stories)", "Done (points)", "% points done"])
start = 4
for j, m in enumerate(meta.MILESTONES):
    r = start + j
    R = f"Backlog!$K$2:$K${last}"
    wsu.append([m["id"], m["title"], f"=COUNTIFS({R},A{r})",
                f"=COUNTIFS({R},A{r},Backlog!$H$2:$H${last},\"Must\")",
                f"=COUNTIFS({R},A{r},Backlog!$H$2:$H${last},\"Should\")",
                f"=COUNTIFS({R},A{r},Backlog!$H$2:$H${last},\"Could\")",
                f"=SUMIFS(Backlog!$J$2:$J${last},{R},A{r})",
                f"=COUNTIFS({R},A{r},Backlog!$N$2:$N${last},\"Done\")",
                f"=SUMIFS(Backlog!$J$2:$J${last},{R},A{r},Backlog!$N$2:$N${last},\"Done\")",
                f"=IF(G{r}=0,0,I{r}/G{r})"])
end = start + len(meta.MILESTONES) - 1
tr = end + 1
wsu.append(["Total", "", f"=SUM(C{start}:C{end})", f"=SUM(D{start}:D{end})", f"=SUM(E{start}:E{end})", f"=SUM(F{start}:F{end})",
            f"=SUM(G{start}:G{end})", f"=SUM(H{start}:H{end})", f"=SUM(I{start}:I{end})", f"=IF(G{tr}=0,0,I{tr}/G{tr})"])
for r in range(start, tr + 1): wsu[f"J{r}"].number_format = "0.0%"
wsu.append([])
mvp_row = tr + 2
wsu.append(["MVP points per sprint (8 sprints)", "", f"=G{start}/8"])
wsu[f"C{mvp_row}"].number_format = "0.0"
wsu.append([])
es = tr + 4
wsu.append(["Epic", "Name", "Stories", "Points", "Done (points)", "% done"])
for k, e in enumerate(EPIC_ORDER):
    r = es + 1 + k
    wsu.append([e, EPIC_NAME[e], f"=COUNTIFS(Backlog!$B$2:$B${last},A{r})", f"=SUMIFS(Backlog!$J$2:$J${last},Backlog!$B$2:$B${last},A{r})",
                f"=SUMIFS(Backlog!$J$2:$J${last},Backlog!$B$2:$B${last},A{r},Backlog!$N$2:$N${last},\"Done\")", f"=IF(D{r}=0,0,E{r}/D{r})"])
    wsu[f"F{r}"].number_format = "0.0%"
ss_row = es + 2 + len(EPIC_ORDER)
wsu.append([])
wsu.append(["Sprint", "", "Stories", "Points", "Done (points)"])
for n in range(0, 25):
    r = ss_row + 1 + n
    wsu.append([f"S{n}", "", f"=COUNTIFS(Backlog!$L$2:$L${last},A{r})", f"=SUMIFS(Backlog!$J$2:$J${last},Backlog!$L$2:$L${last},A{r})",
                f"=SUMIFS(Backlog!$J$2:$J${last},Backlog!$L$2:$L${last},A{r},Backlog!$N$2:$N${last},\"Done\")"])
for r in (3, es, ss_row):
    for c in wsu[r]:
        if c.value is not None: c.font = hdr_font; c.fill = hdr_fill
for row in wsu.iter_rows(min_row=4):
    for c in row:
        if c.font != hdr_font: c.font = body
for r in (tr, mvp_row):
    for c in wsu[r]: c.font = bold
for col, wdt in zip("ABCDEFGHIJ", [12, 44, 9, 9, 9, 9, 9, 13, 13, 12]): wsu.column_dimensions[col].width = wdt

# Personas
wp = wb.create_sheet("Personas")
wp.append(["Name", "Age", "Role", "Summary", "Needs", "Frustrations"])
for p in meta.PERSONAS: wp.append([p["name"], p["age"], p["role"], p["summary"], p["needs"], p["frustrations"]])
for c in wp[1]: c.font = hdr_font; c.fill = hdr_fill
for row in wp.iter_rows(min_row=2):
    for c in row: c.font = body; c.alignment = wrap
for col, wdt in zip("ABCDEF", [12, 6, 26, 60, 40, 40]): wp.column_dimensions[col].width = wdt

# Screens
wsc = wb.create_sheet("Screens")
wsc.append(["Screen", "Title", "Device", "Prototype artboard", "Stories (count)"])
for i, (k, (t, d, f)) in enumerate(meta.SCREENS.items(), start=2):
    wsc.append([k, t, d, f, f'=COUNTIF(Backlog!$M$2:$M${last},"*"&A{i}&"*")'])
for c in wsc[1]: c.font = hdr_font; c.fill = hdr_fill
for row in wsc.iter_rows(min_row=2):
    for c in row: c.font = body
for col, wdt in zip("ABCDE", [8, 32, 10, 30, 14]): wsc.column_dimensions[col].width = wdt

wb.move_sheet("Lists", offset=10)
wb.save("Trailbuddies-backlog.xlsx")
print(len(STORIES), "stories;", dict(cnt), dict(pts))
